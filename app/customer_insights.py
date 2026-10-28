from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def _parse_date(value: Any) -> date | None:
    if value is None:
        return None

    text = str(value).strip()
    if not text or text.lower() in {
        "none", "null", "nan", "n/a", "na", "unknown", "-"
    }:
        return None

    try:
        return datetime.fromisoformat(
            text.replace("Z", "+00:00")
        ).date()
    except ValueError:
        pass

    for fmt in (
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%d-%b-%Y",
        "%d %b %Y",
        "%d.%m.%Y",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    return None


def is_renewal_expiry_question(question: str) -> bool:
    q = question.casefold()

    terms = (
        "expiry", "expire", "expired", "expiring", "renewal",
        "renew", "due date", "due soon", "validity",
        "insurance", "puc", "fitness", "tax upto",
        "document status", "membership end", "membership expiry",
        "renew kar", "renewal kab", "kab expire",
        "समाप्त", "नवीनीकरण", "बीमा", "कब तक वैध",
        "एक्सपायर", "एक्सपायरी",
    )
    return any(term in q for term in terms)


def is_customer_360_question(question: str) -> bool:
    q = question.casefold()

    phrases = (
        "customer 360", "360 summary", "360 view",
        "complete customer summary", "complete account summary",
        "full account summary", "my complete summary",
        "my overall summary", "complete profile",
        "account overview", "my profile summary",
        "meri complete details", "mera complete summary",
        "mere account ka summary", "meri saari details",
        "mere saare details", "full summary",
        "complete summary", "overall summary",
    )
    return any(phrase in q for phrase in phrases)


def is_report_request(question: str) -> bool:
    q = question.casefold()

    report_terms = (
        "generate report", "generate my report", "create report",
        "create my report", "prepare my report", "download report",
        "download my report", "export report", "export my report",
        "make report", "make my report", "report bana",
        "report banao", "report bana do", "meri report",
        "mera report", "my report", "excel report",
        "download excel", "export excel", "download csv",
        "export csv", "excel file bana", "report generate",
        "download my data", "export my data",
    )
    return any(term in q for term in report_terms)


def _expiry_items(
    snapshot: dict[str, Any],
    today: date | None = None,
) -> list[dict[str, Any]]:
    today = today or date.today()
    items: list[dict[str, Any]] = []

    for vehicle_item in snapshot.get("vehicles", []):
        vehicle = vehicle_item.get("vehicle") or {}
        registration = vehicle.get("vehicle_number") or "Unknown vehicle"

        for document in vehicle_item.get("documents", []):
            expiry_text = document.get("expiry_date")
            expiry = _parse_date(expiry_text)

            if expiry is None:
                # Preserve source-labelled expired flags but never invent a date.
                status = str(document.get("status") or "Date unavailable")
                if status.casefold().startswith("expired"):
                    items.append({
                        "vehicle_number": registration,
                        "item_type": "Document",
                        "item": document.get("document") or "Document",
                        "expiry_date": "Date unavailable",
                        "days_left": None,
                        "status": status,
                    })
                continue

            days_left = (expiry - today).days
            if days_left < 0:
                status = f"Expired ({abs(days_left)} days ago)"
            elif days_left == 0:
                status = "Expires today"
            elif days_left <= 30:
                status = f"Expiring in {days_left} days"
            else:
                status = f"Valid ({days_left} days remaining)"

            items.append({
                "vehicle_number": registration,
                "item_type": "Document",
                "item": document.get("document") or "Document",
                "expiry_date": expiry.isoformat(),
                "days_left": days_left,
                "status": status,
            })

        for membership in vehicle_item.get("memberships", []):
            expiry_text = membership.get("end_date")
            expiry = _parse_date(expiry_text)

            if expiry is None:
                continue

            days_left = (expiry - today).days
            if days_left < 0:
                status = f"Expired ({abs(days_left)} days ago)"
            elif days_left == 0:
                status = "Expires today"
            elif days_left <= 30:
                status = f"Expiring in {days_left} days"
            else:
                status = f"Valid ({days_left} days remaining)"

            items.append({
                "vehicle_number": registration,
                "item_type": "Membership",
                "item": membership.get("membership_type") or "Membership",
                "expiry_date": expiry.isoformat(),
                "days_left": days_left,
                "status": status,
            })

    items.sort(
        key=lambda item: (
            item["days_left"] is None,
            item["days_left"] if item["days_left"] is not None else 999999,
            item["vehicle_number"],
        )
    )
    return items


def build_renewal_expiry_answer(
    snapshot: dict[str, Any],
    question: str,
    language: str = "English",
) -> str:
    q = question.casefold()
    lang = (language or "English").strip().casefold()
    items = _expiry_items(snapshot)

    if not items:
        messages = {
            "hindi": "Aapke linked records mein koi readable document ya membership expiry date nahi mili.",
            "hinglish": "Aapke linked records mein koi readable document ya membership expiry date nahi mili.",
        }
        return messages.get(
            lang,
            "No readable document or membership expiry dates were found in your linked records.",
        )

    expired_only = any(word in q for word in (
        "expired", "expire ho chuki", "expire ho gaya",
        "already expired", "समाप्त हो गया", "एक्सपायर हो चुका",
    ))
    upcoming_only = any(word in q for word in (
        "expiring", "upcoming", "next ", "within ",
        "due soon", "agale", "agle", "aane wale",
        "jaldi expire", "within next", "अगले", "आगामी",
    ))

    threshold = 30
    match = re.search(
        r"(?:next|within|agle|agale|aane wale|अगले)\s*(\d{1,3})\s*(?:day|days|din|दिन)",
        q,
    )
    if match:
        threshold = max(0, min(int(match.group(1)), 365))

    dated_items = [item for item in items if item["days_left"] is not None]

    if expired_only:
        selected = [item for item in items
                    if item["days_left"] is not None and item["days_left"] < 0]
        title = "Expired items"
    elif upcoming_only:
        selected = [
            item for item in dated_items
            if 0 <= item["days_left"] <= threshold
        ]
        title = f"Items expiring within {threshold} days"
    else:
        selected = items
        title = "Document and membership expiry overview"

    customer = snapshot.get("customer") or {}
    vehicle_count = snapshot.get("vehicle_count_in_snapshot", 0)

    if not selected:
        if lang == "hindi":
            return f"{title}: aapke linked records mein koi matching item nahi mila. Linked vehicles: {vehicle_count}."
        if lang == "hinglish":
            return f"{title}: aapke linked records mein koi matching item nahi mila. Linked vehicles: {vehicle_count}."
        return f"{title}: no matching items were found in your linked records. Linked vehicles: {vehicle_count}."

    lines = [
        title,
        f"Customer: {customer.get('name') or 'Name unavailable'}",
        f"Linked vehicles in source snapshot: {vehicle_count}",
        f"Matching items: {len(selected)}",
        "",
    ]

    for item in selected[:40]:
        lines.append(
            f"- {item['vehicle_number']} | {item['item_type']}: "
            f"{item['item']} | Expiry: {item['expiry_date']} | "
            f"Status: {item['status']}"
        )

    if len(selected) > 40:
        lines.append(f"...and {len(selected) - 40} more items. Use a report for the complete list.")

    return "\n".join(lines)


def build_customer_360_answer(snapshot: dict[str, Any]) -> str:
    customer = snapshot.get("customer") or {}
    vehicles = snapshot.get("vehicles", [])
    memberships = []
    claim_records = 0
    claim_flag_vehicles = 0
    rc_records = 0

    for item in vehicles:
        memberships.extend(item.get("memberships", []))

        indicator = item.get("claim_indicator") or {}
        if indicator.get("matching_records", 0):
            claim_records += int(indicator.get("matching_records", 0))
        if indicator.get("has_claim_flag"):
            claim_flag_vehicles += 1

        rc_records += len(item.get("rc_statuses", []))

    expiry_items = _expiry_items(snapshot)
    expired = sum(
        1 for item in expiry_items
        if item["days_left"] is not None and item["days_left"] < 0
    )
    upcoming = sum(
        1 for item in expiry_items
        if item["days_left"] is not None
        and 0 <= item["days_left"] <= 30
    )

    lines = [
        "GetRightCover Customer 360 Summary",
        f"Customer: {customer.get('name') or 'Name unavailable'}",
        f"State: {customer.get('state') or 'Not available'}",
        f"Category: {customer.get('category') or 'Not available'}",
        f"Business: {customer.get('business') or 'Not available'}",
        "",
        "Vehicles and memberships",
        f"- Linked vehicles found in snapshot: {snapshot.get('vehicle_count_in_snapshot', len(vehicles))}",
        f"- Membership records: {len(memberships)}",
        "",
        "Renewal and expiry",
        f"- Expired dated items: {expired}",
        f"- Items expiring within 30 days: {upcoming}",
        "",
        "Claims and RC",
        f"- Matching renewal/claim records: {claim_records}",
        f"- Vehicles with a claim flag in the source: {claim_flag_vehicles}",
        f"- RC status records: {rc_records}",
        "",
        "Note: counts reflect available linked source records. Missing or row-level challan data is not inferred.",
    ]
    return "\n".join(lines)


def _style_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="203864")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for column_cells in ws.columns:
        max_length = 0
        column_index = column_cells[0].column
        for cell in column_cells[:250]:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 55))
        ws.column_dimensions[get_column_letter(column_index)].width = max(
            12, min(max_length + 2, 55)
        )


def create_customer_report(snapshot: dict[str, Any]) -> bytes:
    """Create an in-memory XLSX report from one customer's authorized snapshot."""
    customer = snapshot.get("customer") or {}
    vehicles = snapshot.get("vehicles", [])
    expiry_items = _expiry_items(snapshot)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Customer Summary"
    summary.append(["Field", "Value"])
    summary_rows = [
        ("Customer name", customer.get("name")),
        ("State", customer.get("state")),
        ("Category", customer.get("category")),
        ("Business", customer.get("business")),
        ("Vehicles in snapshot", snapshot.get("vehicle_count_in_snapshot", len(vehicles))),
        ("Generated on", date.today().isoformat()),
    ]
    for row in summary_rows:
        summary.append(list(row))

    vehicle_sheet = workbook.create_sheet("Vehicles")
    vehicle_sheet.append([
        "Vehicle Number", "Make", "Model", "State", "Registered At",
        "Insurance Company", "Financier", "Membership Count",
        "Claim Flag", "RC Statuses",
    ])

    memberships_sheet = workbook.create_sheet("Memberships")
    memberships_sheet.append([
        "Vehicle Number", "Membership Type", "Start Date", "End Date",
        "Status", "Expiry Bucket", "Business Partner", "Vehicle Type",
    ])

    expiry_sheet = workbook.create_sheet("Renewal & Expiry")
    expiry_sheet.append([
        "Vehicle Number", "Item Type", "Item", "Expiry Date", "Days Left", "Status",
    ])

    claim_sheet = workbook.create_sheet("Claims & RC")
    claim_sheet.append([
        "Vehicle Number", "Claim Records", "Has Claim Flag", "RC Status",
        "RC Owner", "RC Marked Count",
    ])

    for item in vehicles:
        vehicle = item.get("vehicle") or {}
        reg = vehicle.get("vehicle_number") or "Unknown"
        indicator = item.get("claim_indicator") or {}
        rc_values = item.get("rc_statuses", [])

        vehicle_sheet.append([
            reg, vehicle.get("make"), vehicle.get("model"),
            vehicle.get("state"), vehicle.get("registered_at"),
            vehicle.get("insurance_company"), vehicle.get("financier"),
            len(item.get("memberships", [])),
            bool(indicator.get("has_claim_flag")),
            "; ".join(str(rc.get("status") or "") for rc in rc_values),
        ])

        for membership in item.get("memberships", []):
            memberships_sheet.append([
                reg, membership.get("membership_type"),
                membership.get("start_date"), membership.get("end_date"),
                membership.get("status"), membership.get("expiry_bucket"),
                membership.get("business_partner"), membership.get("vehicle_type"),
            ])

        for rc in rc_values:
            claim_sheet.append([
                reg,
                indicator.get("matching_records", 0),
                bool(indicator.get("has_claim_flag")),
                rc.get("status"), rc.get("owner"), rc.get("marked_count"),
            ])

        if not rc_values:
            claim_sheet.append([
                reg,
                indicator.get("matching_records", 0),
                bool(indicator.get("has_claim_flag")),
                None, None, None,
            ])

    for item in expiry_items:
        expiry_sheet.append([
            item.get("vehicle_number"),
            item.get("item_type"),
            item.get("item"),
            item.get("expiry_date"),
            item.get("days_left"),
            item.get("status"),
        ])

    for ws in workbook.worksheets:
        _style_sheet(ws)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
